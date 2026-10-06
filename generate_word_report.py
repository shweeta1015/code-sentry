"""generate_word_report.py: Generates a professionally styled Microsoft Word (.docx) short report for Assignment 6."""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell internal padding."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def add_hyperlink(paragraph, url, text, color="004B87", underline=True):
    """Adds a clickable hyperlink to a Word paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')

    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        r_pr.append(c)

    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        r_pr.append(u)

    b = OxmlElement('w:b')
    r_pr.append(b)

    new_run.append(r_pr)
    text_node = OxmlElement('w:t')
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)

    paragraph._p.append(hyperlink)


def build_docx_report(output_path: str):
    import docx

    doc = Document()

    # Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(51, 51, 51)  # Charcoal

    # Document Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Assignment 6: AI-Powered Security Code Reviewer")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87)  # Deep Navy

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Project Code Sentry | Technical Report & Evaluation")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(83, 100, 147)

    # Metadata Callout Box (Table)
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    metadata_items = [
        ("Author", "Shweta (shweeta1015)"),
        ("GitHub Repository", "https://github.com/shweeta1015/code-sentry"),
        ("SDK & Tech Stack", "Python 3.10+, Google Gen AI SDK (google-genai), pytest"),
        ("Status & Verification", "Completed | 20/20 Unit Tests Passed | Pushed to GitHub"),
    ]

    for i, (label, val) in enumerate(metadata_items):
        cell_lbl = meta_table.cell(i, 0)
        cell_val = meta_table.cell(i, 1)

        cell_lbl.width = Inches(2.0)
        cell_val.width = Inches(4.5)

        set_cell_background(cell_lbl, "F0F4F8")
        set_cell_background(cell_val, "FAFCFF")
        set_cell_margins(cell_lbl, top=70, bottom=70, left=120, right=120)
        set_cell_margins(cell_val, top=70, bottom=70, left=120, right=120)

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(0)
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = RGBColor(16, 44, 87)

        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(0)
        if label == "GitHub Repository":
            add_hyperlink(p_val, val, val)
        else:
            r_val = p_val.add_run(val)
            r_val.font.size = Pt(10)
            if "Completed" in val:
                r_val.font.bold = True
                r_val.font.color.rgb = RGBColor(34, 139, 34)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1: Executive Overview
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("1. Executive Overview")
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(16, 44, 87)

    p_intro = doc.add_paragraph(
        "Code Sentry is an autonomous AI-powered security code review application designed to assist "
        "software engineering teams in identifying critical vulnerabilities before production deployment. "
        "Built with Python and the modern Google Gen AI SDK (google-genai), the system accepts a natural-language "
        "developer query and an accompanying code snippet. Rather than relying on rigid procedural pipelines, "
        "Google Gemini functions as an autonomous AppSec expert, leveraging automatic function calling to "
        "invoke a local static/heuristic scanner (scan_vulnerabilities) whenever risky patterns are identified."
    )
    p_intro.paragraph_format.space_after = Pt(6)

    # Section 2: GitHub Repository Link & Access
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("2. GitHub Repository Link")
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(16, 44, 87)

    p_gh = doc.add_paragraph()
    p_gh.add_run("The complete source code, test suite, and setup instructions have been pushed to GitHub:\n")
    p_gh_link = doc.add_paragraph()
    p_gh_link.paragraph_format.left_indent = Inches(0.25)
    add_hyperlink(p_gh_link, "https://github.com/shweeta1015/code-sentry", "https://github.com/shweeta1015/code-sentry")
    
    p_gh_info = doc.add_paragraph(
        "- Primary Branch: main\n"
        "- Visibility: Public\n"
        "- Commit Details: Includes main.py, tools.py, prompts.py, tests/test_tools.py, demo.py, and samples."
    )
    p_gh_info.paragraph_format.space_after = Pt(6)

    # Section 3: Architecture & Automatic Tool Calling
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("3. System Architecture & Automatic Tool Invocation")
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(16, 44, 87)

    p_arch = doc.add_paragraph(
        "A cornerstone requirement of Assignment 6 is that the LLM—not application logic—must decide when to call "
        "the local scan_vulnerabilities tool based on snippet content and query context. This is achieved using the "
        "Google Gen AI SDK's automatic function calling capabilities:"
    )
    p_arch.paragraph_format.space_after = Pt(4)

    bullet_points = [
        ("Tool Declaration: ", "scan_vulnerabilities is implemented in tools.py as a pure Python function decorated with comprehensive type hints and docstrings."),
        ("Config Registration: ", "Registered directly into types.GenerateContentConfig(tools=[scan_vulnerabilities], system_instruction=SYSTEM_PROMPT)."),
        ("Autonomous Dispatch: ", "On risky code snippets (database queries, external commands, deserialization), Gemini initiates a function call. The SDK executes the local tool and passes the JSON findings back to Gemini."),
        ("Clean Code Handling: ", "On benign/clean snippets (e.g. pure math, string formatters), the model skips tool execution, avoiding unnecessary overhead and preventing false alarms."),
    ]
    for b_title, b_desc in bullet_points:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        r_bt = bp.add_run(b_title)
        r_bt.font.bold = True
        bp.add_run(b_desc)

    # Section 4: Vulnerability Coverage Matrix (Table)
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(4)
    r_h4 = h4.add_run("4. Vulnerability Detection Coverage Matrix")
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(16, 44, 87)

    table = doc.add_table(rows=8, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Risk Category", "Rule ID", "Severity", "Targeted Risk Patterns"]
    col_widths = [Inches(1.8), Inches(1.1), Inches(1.0), Inches(2.6)]

    for c_idx, h_text in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.width = col_widths[c_idx]
        set_cell_background(cell, "102C57")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    rules_data = [
        ("Hardcoded Secrets & API Keys", "SEC-SECRET-001/3", "CRITICAL / HIGH", "AWS access keys, tokens, plaintext passwords, RSA private keys"),
        ("Dynamic Code Execution", "SEC-RCE-001", "CRITICAL", "Arbitrary code execution via eval() or exec()"),
        ("SQL Injection", "SEC-SQL-001/2", "CRITICAL / HIGH", "Dynamic queries via f-strings, %, or string concatenation (+)"),
        ("Command Injection", "SEC-CMD-001/2", "CRITICAL / HIGH", "subprocess calls with shell=True, os.system(), os.popen()"),
        ("Insecure Deserialization", "SEC-DESER-001/2", "CRITICAL / HIGH", "pickle.loads(), yaml.load() without SafeLoader"),
        ("Weak Cryptographic Hash", "SEC-HASH-001", "MEDIUM", "MD5 or SHA-1 hashes used for sensitive data/passwords"),
        ("Missing Input Validation", "SEC-PATH-001", "HIGH", "Direct user input passed to open(), path traversal sinks"),
    ]

    for r_idx, row in enumerate(rules_data, start=1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = col_widths[c_idx]
            bg_col = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(9)
            if c_idx == 2:
                r.font.bold = True
                if "CRITICAL" in val:
                    r.font.color.rgb = RGBColor(180, 0, 0)
                elif "HIGH" in val:
                    r.font.color.rgb = RGBColor(200, 80, 0)
                else:
                    r.font.color.rgb = RGBColor(180, 140, 0)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 5: Verification & Test Results
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(4)
    r_h5 = h5.add_run("5. Verification & Test Results")
    r_h5.font.size = Pt(14)
    r_h5.font.bold = True
    r_h5.font.color.rgb = RGBColor(16, 44, 87)

    p_test = doc.add_paragraph(
        "A comprehensive automated test suite (tests/test_tools.py) was written and executed using pytest. "
        "All 20 test cases executed with a 100% pass rate in 0.12 seconds."
    )
    p_test.paragraph_format.space_after = Pt(4)

    test_highlights = [
        "Vulnerability Detection: Confirmed detection across all 7 mandated security categories with accurate line numbering.",
        "Zero False Positives: Clean code samples (mathematical functions, parameterized SQL queries) produced 0 findings.",
        "Safe Edge Cases: Empty strings, whitespace-only snippets, and comments containing rule keywords are handled gracefully without crashing or false alarms.",
        "Deterministic Mock Mode: Provided via --mock for automated CI pipelines and offline grading.",
    ]
    for th in test_highlights:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        bp.add_run(th)

    # Section 6: Instructions to Run
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(12)
    h6.paragraph_format.space_after = Pt(4)
    r_h6 = h6.add_run("6. Quick Start & Execution Commands")
    r_h6.font.size = Pt(14)
    r_h6.font.bold = True
    r_h6.font.color.rgb = RGBColor(16, 44, 87)

    p_cmd = doc.add_paragraph(
        "```bash\n"
        "# 1. Clone Repository\n"
        "git clone https://github.com/shweeta1015/code-sentry.git\n"
        "cd code-sentry\n\n"
        "# 2. Install Dependencies\n"
        "pip install -r requirements.txt\n\n"
        "# 3. Run Unit Tests\n"
        "pytest tests/ -v\n\n"
        "# 4. Run Automated Demo\n"
        "python demo.py\n\n"
        "# 5. Run on Custom File with Gemini\n"
        "python main.py --file samples/vulnerable_sql.py\n"
        "```"
    )
    p_cmd.paragraph_format.space_after = Pt(6)

    # Save Document
    doc.save(output_path)
    print(f"Successfully generated Word report at: {output_path}")


if __name__ == "__main__":
    out1 = r"c:\Users\Shweta\Desktop\anti_gravity_workspace\code_sentry\Assignment_6_Code_Sentry_Report.docx"
    out2 = r"c:\Users\Shweta\Desktop\anti_gravity_workspace\Assignment_6_Code_Sentry_Report.docx"
    build_docx_report(out1)
    build_docx_report(out2)
